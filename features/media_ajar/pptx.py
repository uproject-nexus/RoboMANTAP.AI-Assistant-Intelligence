"""Media Ajar PPTX renderer. Stable rendering logic preserved."""
from __future__ import annotations
import io
def build_media_ajar_pptx(storyboard: dict, config: dict, logo_path: str | None = None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor as PptRGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches as PptInches, Pt as PptPt

    prs = Presentation()
    prs.slide_width = PptInches(13.333333)
    prs.slide_height = PptInches(7.5)
    NAVY = PptRGBColor(7, 15, 35)
    NAVY2 = PptRGBColor(15, 23, 42)
    EMERALD = PptRGBColor(16, 185, 129)
    MINT = PptRGBColor(167, 243, 208)
    WHITE = PptRGBColor(248, 250, 252)
    SLATE = PptRGBColor(148, 163, 184)
    GOLD = PptRGBColor(245, 158, 11)
    BLUE = PptRGBColor(59, 130, 246)

    def text_box(slide, value, x, y, w, h, size=20, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
        shape = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
        frame = shape.text_frame
        frame.clear(); frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.text = str(value or "")
        paragraph.alignment = align
        run = paragraph.runs[0]
        run.font.size = PptPt(size); run.font.bold = bold; run.font.color.rgb = color
        return shape

    def background(slide, color):
        fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = color

    def card(slide, value, x, y, w, h, accent=EMERALD, size=14):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = NAVY2; shape.line.color.rgb = PptRGBColor(51,65,85)
        strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(x), PptInches(y), PptInches(0.08), PptInches(h))
        strip.fill.solid(); strip.fill.fore_color.rgb = accent; strip.line.fill.background()
        text_box(slide, value, x+0.22, y+0.15, w-0.4, h-0.25, size, False, WHITE)

    for data in storyboard.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide_type = str(data.get("type", "concept")).lower()
        title = data.get("title", "RoboMANTAP Media")
        if slide_type == "cover":
            background(slide, NAVY)
            text_box(slide, "U.PROJECT NEXUS • ROBOMANTAP", 0.7, 0.7, 7.5, 0.3, 12, True, MINT)
            text_box(slide, title, 0.7, 1.7, 10.4, 1.35, 34, True, WHITE)
            text_box(slide, data.get("subtitle", ""), 0.72, 3.15, 9.8, 0.75, 18, False, SLATE)
            pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(0.72), PptInches(4.3), PptInches(2.7), PptInches(0.52))
            pill.fill.solid(); pill.fill.fore_color.rgb = EMERALD; pill.line.fill.background()
            text_box(slide, f"{config.get('kelas','')} • {config.get('mapel','')}", 0.86, 4.43, 2.45, 0.2, 10, True, NAVY)
        else:
            background(slide, NAVY)
            text_box(slide, "ROBOMANTAP • MEDIA STUDIO", 0.55, 0.25, 4.8, 0.25, 10, True, MINT)
            text_box(slide, title, 0.55, 0.75, 11.9, 0.72, 25, True, WHITE)
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.55), PptInches(1.52), PptInches(1.1), PptInches(0.05))
            line.fill.solid(); line.fill.fore_color.rgb = EMERALD; line.line.fill.background()
            visual_type = str(data.get("visual_type", "concept")).lower()
            visual = str(data.get("visual_content", ""))
            body = [str(x) for x in (data.get("body") or []) if str(x).strip()]
            if visual_type in {"diagram", "flow", "process"}:
                steps = body[:5] or [visual]
                width = 11.7 / max(1, len(steps))
                for i, item in enumerate(steps):
                    card(slide, item, 0.7 + i*width, 2.15, width-0.22, 2.0, EMERALD if i % 2 == 0 else GOLD, 13)
            elif visual_type == "table":
                rows = [r for r in visual.splitlines() if r.strip()] or body
                y = 2.05
                for row in rows[:6]:
                    parts = [p.strip() for p in row.split("|")]
                    if len(parts) <= 1:
                        card(slide, row, 0.8, y, 11.7, 0.6, EMERALD, 13)
                    else:
                        cell_w = 11.7 / min(3, len(parts))
                        for c, value in enumerate(parts[:3]):
                            card(slide, value, 0.8 + c*cell_w, y, cell_w-0.12, 0.6, EMERALD if c == 0 else BLUE, 12)
                    y += 0.74
            elif visual_type in {"activity", "discussion", "quiz"}:
                card(slide, visual or (body[0] if body else "Diskusikan dengan kelompok."), 0.75, 2.0, 11.7, 1.55, GOLD, 18)
                for i, item in enumerate(body[1:5]):
                    card(slide, item, 0.75 + (i % 2)*5.95, 3.9 + (i//2)*1.0, 5.65, 0.8, EMERALD, 13)
            else:
                text_box(slide, visual, 0.8, 2.0, 11.55, 1.25, 23, True, MINT)
                y = 3.35
                for i, item in enumerate(body[:5]):
                    card(slide, item, 0.8, y, 11.55, 0.56, EMERALD if i % 2 == 0 else BLUE, 12)
                    y += 0.67
            source = data.get("source_locator", "")
            if source:
                text_box(slide, f"Sumber: {source}", 9.1, 7.02, 3.6, 0.2, 8, False, SLATE, PP_ALIGN.RIGHT)
        try:
            slide.notes_slide.notes_text_frame.text = data.get("speaker_notes", "") or ""
        except Exception:
            pass

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
