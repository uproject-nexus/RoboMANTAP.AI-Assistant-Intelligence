from __future__ import annotations

import io, os, re, json, math, html, tempfile
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from typing import Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from xml.sax.saxutils import escape
from infrastructure.ai.service import call_gemini_with_rotation, clean_json_text
from infrastructure.ai.material import _call_gemini_contents
try:
    from google.genai import types
except Exception:
    types = None

FORM_ORDER = ("PG", "Isian", "Uraian")
COGNITIVE_LEVELS = ("C1", "C2", "C3", "C4", "C5", "C6")

from .common import contains_arabic, apply_arabic_paragraph_style, apply_arabic_run_style, add_omml_fraction, add_omml_matrix, append_text_with_fractions
from .generator import _normalize_generated_question, render_math_diagram

def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))

def _add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14 if level == 1 else 11)
    r.font.color.rgb = RGBColor(6, 78, 59)
    return p

def build_bank_soal_docx(
    blueprint: dict,
    questions: list[dict],
    *,
    jenjang: str,
    mapel: str,
    kelas: str = "",
    variants: int = 1,
    include_answer_key: bool = True,
    include_blueprint_map: bool = True,
    logo_path: str | None = None,
) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    # Header / identity.
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.9)
    c0, c1 = table.rows[0].cells
    c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    if logo_path and os.path.exists(logo_path):
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c0.paragraphs[0].add_run().add_picture(logo_path, width=Inches(0.9))
    else:
        r = c0.paragraphs[0].add_run("UPN")
        r.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = RGBColor(6, 78, 59)

    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ROBO MANTAP • BANK SOAL UJIAN\n")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(6, 78, 59)
    r2 = p.add_run(f"{mapel or 'Mata Pelajaran'} • {jenjang} {kelas}".strip())
    r2.bold = True
    r2.font.size = Pt(10)

    meta = doc.add_table(rows=0, cols=2)
    meta.autofit = False
    metadata = blueprint.get("metadata", {})
    meta_rows = [
        ("Asesmen", metadata.get("assessment") or "-"),
        ("Tahun Pelajaran", metadata.get("school_year") or "-"),
        ("Blueprint", f"{len(blueprint.get('blueprints', []))} item"),
        ("Variasi / Blueprint", str(variants)),
        ("Total Soal", str(len(questions))),
    ]
    for label, value in meta_rows:
        cells = meta.add_row().cells
        cells[0].width = Inches(1.65)
        cells[1].width = Inches(5.35)
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].paragraphs[0].add_run(str(value))
        _set_cell_shading(cells[0], "EAF7F1")
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_border(
                c,
                top={"val": "single", "sz": 4, "color": "D1D5DB"},
                bottom={"val": "single", "sz": 4, "color": "D1D5DB"},
                left={"val": "single", "sz": 4, "color": "D1D5DB"},
                right={"val": "single", "sz": 4, "color": "D1D5DB"},
            )

    if include_blueprint_map:
        _add_heading(doc, "Peta Blueprint", 1)
        t = doc.add_table(rows=1, cols=6)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["ID", "Bab", "ATP", "Indikator", "Level", "Bentuk / Nomor"]
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = h
            _set_cell_shading(cell, "064E3B")
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8)
        for bp in blueprint.get("blueprints", []):
            forms_text = " • ".join(
                f"{form}: {', '.join(map(str, bp.get('forms', {}).get(form, [])))}"
                for form in FORM_ORDER if bp.get("forms", {}).get(form)
            )
            cells = t.add_row().cells
            vals = [
                bp["id"], bp.get("chapter", ""), bp.get("atp", ""),
                bp.get("indicator", ""), bp.get("cognitive_level") or "—", forms_text
            ]
            for i, val in enumerate(vals):
                cells[i].text = str(val)
                for run in cells[i].paragraphs[0].runs:
                    run.font.size = Pt(7.5)
        doc.add_page_break()

    # ------------------------------------------------------------------
    # Final document layout: one complete package per Variant.
    # IMPORTANT: never mix questions from different variants on the same
    # package/page. Each variant is rendered as:
    #   VARIAN N
    #   A. PILIHAN GANDA
    #   B. ISIAN SINGKAT
    #   C. URAIAN
    # followed by a page break before the next variant.
    # ------------------------------------------------------------------
    # Render the configured number of packages explicitly. This prevents a
    # missing/empty variant from being silently merged into another variant.
    requested_variants = max(1, int(variants or blueprint.get("variants_per_blueprint", 1) or 1))
    variant_values = list(range(1, requested_variants + 1))

    form_titles = {
        "PG": "A. PILIHAN GANDA",
        "Isian": "B. ISIAN SINGKAT",
        "Uraian": "C. URAIAN",
    }

    # Build explicit buckets first. This avoids relying on the incoming
    # question order and guarantees V1 is completed before V2 starts.
    variant_buckets = {
        variant: {
            form: [
                q for q in questions
                if int(q.get("variant", 0)) == variant
                and q.get("question_type") == form
            ]
            for form in FORM_ORDER
        }
        for variant in variant_values
    }

    for variant_index, variant in enumerate(variant_values):
        # Every variant starts on a fresh page. This makes each variant a
        # genuinely separate package when the DOCX is printed or distributed.
        if variant_index > 0:
            doc.add_page_break()

        # Variant banner.
        banner = doc.add_table(rows=1, cols=1)
        banner.alignment = WD_TABLE_ALIGNMENT.CENTER
        banner.autofit = False
        banner.columns[0].width = Inches(6.9)
        bc = banner.cell(0, 0)
        bc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(bc, "064E3B")
        _set_cell_border(
            bc,
            top={"val": "single", "sz": 8, "color": "064E3B"},
            bottom={"val": "single", "sz": 8, "color": "064E3B"},
            left={"val": "single", "sz": 8, "color": "064E3B"},
            right={"val": "single", "sz": 8, "color": "064E3B"},
        )
        bp = bc.paragraphs[0]
        bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bp.paragraph_format.space_before = Pt(7)
        bp.paragraph_format.space_after = Pt(7)
        br = bp.add_run(f"VARIAN {variant}")
        br.bold = True
        br.font.size = Pt(16)
        br.font.color.rgb = RGBColor(255, 255, 255)

        # Small package label makes the purpose unambiguous in print/preview.
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(10)
        cr = cp.add_run("PAKET SOAL • SEMUA BENTUK SOAL DALAM VARIAN INI")
        cr.bold = True
        cr.font.size = Pt(7.5)
        cr.font.color.rgb = RGBColor(107, 114, 128)

        for form in FORM_ORDER:
            form_questions = variant_buckets.get(variant, {}).get(form, [])

            # Always keep the three planned section labels in the same order.
            # If a blueprint has no slot for a form, show a compact note rather
            # than silently merging the next form into the current section.
            _add_heading(doc, form_titles[form], 1)

            if not form_questions:
                empty = doc.add_paragraph()
                empty.paragraph_format.left_indent = Inches(0.18)
                empty.paragraph_format.space_after = Pt(7)
                er = empty.add_run("Tidak ada soal untuk bentuk ini pada varian ini.")
                er.italic = True
                er.font.size = Pt(8.5)
                er.font.color.rgb = RGBColor(107, 114, 128)
                continue

            for idx, q in enumerate(form_questions, start=1):
                # The printed number is local to Variant + Bentuk.
                # Traceability stores the same number together with V/form.
                q["_document_number"] = idx

                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(f"{idx}. ")
                r.bold = True
                r.font.size = Pt(10)

                append_text_with_fractions(p, q.get("question", ""))

                diagram_bytes = render_math_diagram(q.get("diagram")) if q.get("diagram") else None
                if diagram_bytes:
                    dp = doc.add_paragraph()
                    dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    dp.paragraph_format.space_before = Pt(3)
                    dp.paragraph_format.space_after = Pt(5)
                    dp.add_run().add_picture(io.BytesIO(diagram_bytes), width=Inches(0.85))
                    if q.get("diagram", {}).get("caption"):
                        cp = doc.add_paragraph()
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cr = cp.add_run(q["diagram"]["caption"])
                        cr.italic = True
                        cr.font.size = Pt(8)
                        cr.font.color.rgb = RGBColor(107, 114, 128)

                if form == "PG":
                    for opt in q.get("options", []):
                        po = doc.add_paragraph()
                        po.paragraph_format.left_indent = Inches(0.22)
                        po.paragraph_format.space_after = Pt(1)
                        append_text_with_fractions(po, opt)

                if include_answer_key:
                    pa = doc.add_paragraph()
                    pa.paragraph_format.left_indent = Inches(0.22)
                    ra = pa.add_run("Kunci: ")
                    ra.bold = True
                    ra.font.size = Pt(8.5)
 
                    append_text_with_fractions(pa, q.get("correct_answer", ""), color_rgb=RGBColor(5, 150, 105))

                    ps = doc.add_paragraph()
                    ps.paragraph_format.left_indent = Inches(0.22)
                    rs = ps.add_run("Pembahasan: ")
                    rs.bold = True
                    rs.font.size = Pt(8.5)

                    append_text_with_fractions(ps, q.get("solution_basis", ""))
                    ps.paragraph_format.space_after = Pt(7)

    # ------------------------------------------------------------------
    # Blueprint traceability appendix. The rows deliberately follow the same
    # human-readable order as the document: V1 PG -> Isian -> Uraian, then V2...
    # ------------------------------------------------------------------
    if include_blueprint_map:
        doc.add_page_break()
        _add_heading(doc, "Lampiran • Traceability Bank Soal", 1)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run("Cara membaca: ")
        r.bold = True
        r.font.size = Pt(8.5)
        p.add_run(
            "No. Soal mengikuti nomor pada bagian Variant dan Bentuk. "
            "Variant membedakan paket soal, sedangkan Blueprint dan No. Kisi "
            "menunjukkan sumber kisi-kisi yang menjadi dasar soal."
        ).font.size = Pt(8.5)

        trace = doc.add_table(rows=1, cols=8)
        trace.alignment = WD_TABLE_ALIGNMENT.CENTER
        trace.autofit = False
        headers = ["Varian", "No. Soal", "Blueprint", "Level", "Bentuk", "No. Kisi", "QA", "Kode"]
        widths = [0.58, 0.62, 0.72, 0.55, 0.68, 0.62, 0.52, 1.28]
        for i, h in enumerate(headers):
            cell = trace.rows[0].cells[i]
            cell.width = Inches(widths[i])
            cell.text = h
            _set_cell_shading(cell, "064E3B")
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(7.5)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        form_rank = {form: idx for idx, form in enumerate(FORM_ORDER)}
        trace_questions = sorted(
            questions,
            key=lambda q: (
                int(q.get("variant", 0)),
                form_rank.get(q.get("question_type"), 99),
                int(q.get("_document_number", 0)),
                q.get("blueprint_id", ""),
                int(q.get("source_number", 0)),
            ),
        )
        for q in trace_questions:
            cells = trace.add_row().cells
            variant = int(q.get("variant", 0))
            form = q.get("question_type", "")
            no_soal = q.get("_document_number", "")
            bp_id = q.get("blueprint_id", "")
            source_no = q.get("source_number", "")
            form_code = {"PG": "PG", "Isian": "IS", "Uraian": "UR"}.get(form, re.sub(r"[^A-Za-z0-9]+", "", str(form)).upper()[:3])
            code = f"{bp_id}-V{variant}-{form_code}-{no_soal:02d}" if isinstance(no_soal, int) else f"{bp_id}-V{variant}-{form_code}-{no_soal}"
            bp_lookup = next((bp for bp in blueprint.get("blueprints", []) if bp.get("id") == bp_id), {})
            vals = [
                f"V{variant}", no_soal, bp_id, bp_lookup.get("cognitive_level") or "—", form, source_no,
                "PASS" if q.get("_qa_status") == "pass" else "REVIEW",
                code,
            ]
            for i, val in enumerate(vals):
                cells[i].width = Inches(widths[i])
                cells[i].text = str(val)
                cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for run in cells[i].paragraphs[0].runs:
                    run.font.size = Pt(7)
                if i == 5:
                    for run in cells[i].paragraphs[0].runs:
                        run.bold = True
                        run.font.color.rgb = (RGBColor(5, 150, 105) if q.get("_qa_status") == "pass" else RGBColor(180, 83, 9))

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
