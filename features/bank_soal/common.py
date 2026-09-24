from __future__ import annotations

import io, os, re, json, math, html, tempfile
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from typing import Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from xml.sax.saxutils import escape
from infrastructure.ai.service import call_gemini_with_rotation, clean_json_text
from infrastructure.ai.material import _call_gemini_contents
try:
    from google.genai import types
except Exception:
    types = None

FORM_ORDER = ("PG", "Isian", "Uraian")
COGNITIVE_LEVELS = ("C1", "C2", "C3", "C4", "C5", "C6")



def _normalize_cognitive_level(value: Any) -> str:
    """Normalize explicit Bloom/cognitive labels such as C4, C-4, C4 (Analisis).

    Multiple levels are preserved as a slash-separated value so the source
    blueprint is not silently rewritten.
    """
    text = _clean(value).upper()
    if not text:
        return ""
    found = []
    for n in re.findall(r"\bC\s*[-–]?\s*([1-6])\b", text):
        label = f"C{n}"
        if label not in found:
            found.append(label)
    return "/".join(found)

def _cognitive_levels(value: Any) -> list[str]:
    normalized = _normalize_cognitive_level(value)
    return [x for x in normalized.split("/") if x in COGNITIVE_LEVELS]

def _find_cognitive_column(rows: list[list[str]], search_upto: int) -> int | None:
    """Find an explicit cognitive-level column without guessing from content."""
    labels = (
        "LEVEL KOGNITIF", "TINGKAT KOGNITIF", "KOGNITIF",
        "TAKSONOMI BLOOM", "TAKSONOMI", "LEVEL KOGNISI", "LEVEL SOAL",
        "TINGKAT KESULITAN", "LEVEL KESULITAN", "LEVEL",
    )
    for row in rows[: max(1, search_upto + 1)]:
        for idx, cell in enumerate(row):
            value = _clean(cell).upper()
            if any(label in value for label in labels):
                return idx
    return None

def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()

def clean_math_string(text: str) -> str:
    """Same math-cleaning logic used by the existing Quiz DOCX renderer.

    Bank Soal intentionally shares this renderer so fractions, matrices,
    inverses, composition symbols and escaped newlines behave the same way.
    """
    if not text:
        return ""
    text = str(text)

    # Literal and actual newlines must never leak into the Word document.
    text = text.replace("\\r\\n", " ").replace("\\n", " ").replace("\\r", " ")
    text = text.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")

    replacements = {
        r"\rightarrow": "→", r"\to": "→", r"\Rightarrow": "⇒",
        r"\leftarrow": "←", r"\leftrightarrow": "↔",
        r"\Longleftrightarrow": "⇔", r"\longleftrightarrow": "↔",
        r"\Longleftarrow": "⇐", r"\Longrightarrow": "⇒",
        r"\implies": "⇒", r"\impliedby": "⇐", r"\iff": "⇔",
        r"\circ": "∘", r"\circl": "∘",
        r"\times": "×", r"\cdot": "·", r"\div": "÷", r"\neq": "≠",
        r"\leq": "≤", r"\geq": "≥", r"\le": "≤", r"\ge": "≥",
        r"\pm": "±", r"\mp": "∓", r"\infty": "∞", r"\pi": "π",
        r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\delta": "δ",
        r"\theta": "θ", r"\lambda": "λ", r"\mu": "μ", r"\sigma": "σ",
        r"\in": "∈", r"\notin": "∉", r"\forall": "∀", r"\exists": "∃",
        r"\emptyset": "∅", r"\angle": "∠", r"\perp": "⊥", r"\parallel": "∥",
        r"\approx": "≈", r"\equiv": "≡", r"\propto": "∝",
        r"\sum": "Σ", r"\prod": "Π", r"\int": "∫", r"\partial": "∂", r"\nabla": "∇",
        r"\Delta": "Δ", r"\Omega": "Ω", r"\Gamma": "Γ", r"\Lambda": "Λ",
        r"\Sigma": "Σ", r"\Phi": "Φ", r"\degree": "°",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"(?<![A-Za-z])circl(?![A-Za-z])", "∘", text)
    text = re.sub(r"\\left\b\s*[\(\[\{\.\|]?", "(", text)
    text = re.sub(r"\\right\b\s*[\)\]\}\.\|]?", ")", text)
    text = re.sub(r"\\(?:dots|cdots|ldots)", "…", text)
    text = re.sub(r"\\sqrt\{([^}]+)\}", r"√(\1)", text)
    text = re.sub(r"\\sqrt\s*([a-zA-Z0-9_]+)", r"√\1", text)
    sup_map = str.maketrans("0123456789+-=()nxyi", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿˣʸⁱ")
    sub_map = str.maketrans("0123456789+-=()nixy", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₙᵢₓᵧ")
    text = re.sub(r"\^\{([^}]+)\}|\^([\-0-9a-zA-Z])", lambda m: (m.group(1) or m.group(2)).translate(sup_map), text)
    text = re.sub(r"_\{([^}]+)\}|_([0-9a-zA-Z])", lambda m: (m.group(1) or m.group(2)).translate(sub_map), text)
    text = text.replace("$", "")
    text = text.replace("left(", "(").replace("right)", ")").replace("dots", "…")
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\\([a-zA-Z]+)", r"\1", text).replace("\\", "")
    return re.sub(r"\s+", " ", text).strip()

def contains_arabic(text: str) -> bool:
    return bool(re.search(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF]", str(text or "")))

def apply_arabic_paragraph_style(paragraph):
    try:
        ppr = paragraph._p.get_or_add_pPr()
        bidi = parse_xml(r'<w:bidi xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="1"/>')
        ppr.append(bidi)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    except Exception:
        pass

def apply_arabic_run_style(run):
    try:
        run.font.name = "Traditional Arabic"
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
            rpr.append(rfonts)
        rfonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs', 'Traditional Arabic')
        rfonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii', 'Aptos')
    except Exception:
        pass

def add_omml_fraction(paragraph, num_text: str, den_text: str):
    num_clean = clean_math_string(num_text)
    den_clean = clean_math_string(den_text)
    omml_xml = (
        f'<m:oMath {nsdecls("m")}>'
        f'<m:f>'
        f'<m:num><m:r><m:t>{escape(num_clean)}</m:t></m:r></m:num>'
        f'<m:den><m:r><m:t>{escape(den_clean)}</m:t></m:r></m:den>'
        f'</m:f>'
        f'</m:oMath>'
    )
    paragraph._p.append(parse_xml(omml_xml))

def add_omml_matrix(paragraph, matrix_type: str, content: str):
    beg_chr, end_chr = "(", ")"
    if matrix_type == "bmatrix":
        beg_chr, end_chr = "[", "]"
    elif matrix_type in ["vmatrix", "Vmatrix"]:
        beg_chr, end_chr = "|", "|"
    elif matrix_type == "matrix":
        beg_chr, end_chr = "", ""
    rows = [r.strip() for r in re.split(r"\\\\|\\cr", content) if r.strip()]
    matrix_xml_rows = []
    for row in rows:
        cols = [c.strip() for c in row.split("&")]
        cols_xml = []
        for col in cols:
            cleaned = html.escape(clean_math_string(col))
            cols_xml.append(f'<m:e><m:r><m:t>{cleaned}</m:t></m:r></m:e>')
        matrix_xml_rows.append(f'<m:mr>{"".join(cols_xml)}</m:mr>')
    inner_matrix = f'<m:m>{"".join(matrix_xml_rows)}</m:m>'
    if beg_chr or end_chr:
        omml_xml = (
            f'<m:oMath {nsdecls("m")}><m:d><m:dPr>'
            f'<m:begChr m:val="{escape(beg_chr)}"/><m:endChr m:val="{escape(end_chr)}"/>'
            f'</m:dPr><m:e>{inner_matrix}</m:e></m:d></m:oMath>'
        )
    else:
        omml_xml = f'<m:oMath {nsdecls("m")}>{inner_matrix}</m:oMath>'
    paragraph._p.append(parse_xml(omml_xml))

def append_text_with_fractions(paragraph, text: str, is_bold: bool = False, color_rgb: RGBColor = None):
    """Shared Quiz DOCX math renderer: plain text + vertical fractions + matrices."""
    if not text:
        return
    text = str(text)
    math_pattern = re.compile(
        r"\\begin\{(?P<mtype>[pbvV]?matrix)\}(?P<mcontent>.*?)\\end\{(?P=mtype)\}|\\(?:f|tf)rac\{(?P<num>[^}]+)\}\{(?P<den>[^}]+)\}",
        re.DOTALL,
    )
    last_idx = 0
    for match in math_pattern.finditer(text):
        start, end = match.span()
        if start > last_idx:
            plain_part = clean_math_string(text[last_idx:start])
            if plain_part:
                run = paragraph.add_run(plain_part + " ")
                run.bold = is_bold
                if contains_arabic(plain_part):
                    apply_arabic_paragraph_style(paragraph)
                    apply_arabic_run_style(run)
                if color_rgb:
                    run.font.color.rgb = color_rgb
        if match.group("mtype"):
            add_omml_matrix(paragraph, match.group("mtype"), match.group("mcontent"))
        elif match.group("num"):
            add_omml_fraction(paragraph, match.group("num"), match.group("den"))
        run_space = paragraph.add_run(" ")
        run_space.bold = is_bold
        last_idx = end
    if last_idx < len(text):
        plain_part = clean_math_string(text[last_idx:])
        if plain_part:
            run = paragraph.add_run(plain_part)
            run.bold = is_bold
            if contains_arabic(plain_part):
                apply_arabic_paragraph_style(paragraph)
                apply_arabic_run_style(run)
            if color_rgb:
                run.font.color.rgb = color_rgb

def _cell_clean(value: Any) -> str:
    text = str(value or "").replace("\xa0", " ")
    # Merged cells in Word can expose duplicated line fragments.
    parts = [_clean(p) for p in text.splitlines() if _clean(p)]
    if not parts:
        return ""
    return parts[0] if all(p == parts[0] for p in parts) else " ".join(parts)

def _parse_question_numbers(value: str) -> list[int]:
    """Parse 1, 1–2, 1-3 and similar blueprint numbering."""
    raw = _clean(value).replace("—", "-").replace("–", "-")
    if not raw:
        return []
    numbers: list[int] = []
    for chunk in re.split(r"[,;/]+", raw):
        chunk = chunk.strip()
        if not chunk:
            continue
        match = re.fullmatch(r"(\d+)\s*-\s*(\d+)", chunk)
        if match:
            start, end = int(match.group(1)), int(match.group(2))
            if end >= start and end - start <= 200:
                numbers.extend(range(start, end + 1))
            continue
        for n in re.findall(r"\d+", chunk):
            numbers.append(int(n))
    return list(dict.fromkeys(numbers))

def _is_chapter_row(cells: list[str]) -> bool:
    first = _clean(cells[0]) if cells else ""
    if not first.upper().startswith("BAB "):
        return False
    nonempty = [_clean(x).upper() for x in cells[:3] if _clean(x)]
    return len(nonempty) >= 2 and len(set(nonempty)) == 1

def _looks_like_header(cells: list[str]) -> bool:
    joined = " ".join(_clean(x).upper() for x in cells)
    return "ATP" in joined and "INDIKATOR" in joined

def _looks_like_form_header(cells: list[str]) -> bool:
    return any(_clean(x).upper() in FORM_ORDER for x in cells)
