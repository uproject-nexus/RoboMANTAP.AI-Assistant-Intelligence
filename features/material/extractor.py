"""Material extraction primitives for DOCX, PPT/PPTX, PDF and images."""
from __future__ import annotations
import hashlib, io, os, subprocess, tempfile
from infrastructure.ai.service import vision_describe

SUPPORTED_MATERIAL_EXTENSIONS={".docx":"application/vnd.openxmlformats-officedocument.wordprocessingml.document",".pdf":"application/pdf",".pptx":"application/vnd.openxmlformats-officedocument.presentationml.presentation",".ppt":"application/vnd.ms-powerpoint",".png":"image/png",".jpg":"image/jpeg",".jpeg":"image/jpeg",".webp":"image/webp"}

def trim_text(value,max_chars=30000):
    value=str(value or "").strip()
    return value if len(value)<=max_chars else value[:max_chars]+"\n...[dipotong untuk menjaga konteks AI]"

def extract_docx_content(raw_bytes):
    from docx import Document
    doc=Document(io.BytesIO(raw_bytes)); blocks=[]
    for p in doc.paragraphs:
        if p.text.strip(): blocks.append(p.text.strip())
    for idx,t in enumerate(doc.tables,1):
        rows=[" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        if rows: blocks.append(f"[TABEL {idx}]\n"+"\n".join(rows))
    return trim_text("\n\n".join(blocks),40000)

def convert_legacy_ppt_to_pptx(raw_bytes):
    workdir=tempfile.mkdtemp(prefix="robomantap_ppt_"); src=os.path.join(workdir,"source.ppt")
    pathlib.Path(src).write_bytes(raw_bytes)
    try:
        subprocess.run(["soffice","--headless","--convert-to","pptx","--outdir",workdir,src],stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=45,check=False)
        dst=os.path.join(workdir,"source.pptx")
        return pathlib.Path(dst).read_bytes() if os.path.exists(dst) else None
    except Exception:
        return None

def extract_pptx_content(raw_bytes):
    from pptx import Presentation
    prs=Presentation(io.BytesIO(raw_bytes)); slides=[]; images=[]
    for slide_no,slide in enumerate(prs.slides,1):
        texts=[]
        for shape in slide.shapes:
            if getattr(shape,"has_text_frame",False) and shape.text.strip(): texts.append(shape.text.strip())
            try:
                if getattr(shape,"shape_type",None)==13 and getattr(shape,"image",None): images.append({"slide":slide_no,"mime_type":shape.image.content_type or "image/png","bytes":shape.image.blob})
            except Exception: pass
        notes=""
        try: notes=slide.notes_slide.notes_text_frame.text.strip()
        except Exception: pass
        slides.append({"slide":slide_no,"text":"\n".join(texts),"notes":notes})
    return {"slides":slides,"images":images}

def extract_pdf_content(raw_bytes):
    import fitz
    pdf=fitz.open(stream=raw_bytes,filetype="pdf"); pages=[]; image_pages=[]
    for page_no,page in enumerate(pdf,1):
        txt=page.get_text("text").strip(); pages.append({"page":page_no,"text":trim_text(txt,7000)})
        if len(txt)<80 and len(image_pages)<6:
            pix=page.get_pixmap(matrix=fitz.Matrix(1.35,1.35),alpha=False); image_pages.append({"page":page_no,"mime_type":"image/png","bytes":pix.tobytes("png")})
    return {"pages":pages,"images":image_pages}

def extract_image_content(raw_bytes,mime_type,source_label):
    return vision_describe(raw_bytes,mime_type,source_label)
