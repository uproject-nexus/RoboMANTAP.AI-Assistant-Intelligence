"""Source-material extraction and AI grounding helpers."""
from __future__ import annotations
import io, re, json, os, subprocess, tempfile
from infrastructure.ai.service import call_gemini_with_rotation, clean_json_text
def normalize_material_trigger(value: str) -> bool:
    normalized = re.sub(r"\s+", " ", str(value or "").strip().lower())
    return normalized == "materi dilampirkan"


def _trim_text(value: str, max_chars: int = 30000) -> str:
    value = str(value or "").strip()
    return value if len(value) <= max_chars else value[:max_chars] + "\n...[dipotong untuk menjaga konteks AI]"


def _extract_docx_content(raw_bytes: bytes) -> str:
    from docx import Document as _Document
    doc = _Document(io.BytesIO(raw_bytes))
    blocks = []
    for paragraph in doc.paragraphs:
        text_value = paragraph.text.strip()
        if text_value:
            blocks.append(text_value)
    for idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        if rows:
            blocks.append(f"[TABEL {idx}]\n" + "\n".join(rows))
    return _trim_text("\n\n".join(blocks), 40000)


def _convert_legacy_ppt_to_pptx(raw_bytes: bytes):
    workdir = tempfile.mkdtemp(prefix="robomantap_ppt_")
    source_path = os.path.join(workdir, "source.ppt")
    with open(source_path, "wb") as handle:
        handle.write(raw_bytes)
    try:
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pptx", "--outdir", workdir, source_path],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=45,
            check=False,
        )
        output_path = os.path.join(workdir, "source.pptx")
        if result.returncode == 0 and os.path.exists(output_path):
            with open(output_path, "rb") as handle:
                return handle.read()
    except Exception:
        pass
    return None


def _extract_pptx_content(raw_bytes: bytes):
    from pptx import Presentation as _Presentation
    prs = _Presentation(io.BytesIO(raw_bytes))
    slides = []
    images = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text_value = shape.text.strip()
                if text_value:
                    texts.append(text_value)
            try:
                if getattr(shape, "shape_type", None) == 13 and getattr(shape, "image", None):
                    images.append({
                        "slide": slide_no,
                        "mime_type": shape.image.content_type or "image/png",
                        "bytes": shape.image.blob,
                    })
            except Exception:
                pass
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass
        slides.append({"slide": slide_no, "text": "\n".join(texts), "notes": notes})
    return {"slides": slides, "images": images}


def _extract_pdf_content(raw_bytes: bytes):
    import fitz
    pdf = fitz.open(stream=raw_bytes, filetype="pdf")
    pages = []
    image_pages = []
    for page_no, page in enumerate(pdf, start=1):
        text_value = page.get_text("text").strip()
        pages.append({"page": page_no, "text": _trim_text(text_value, 7000)})
        if len(text_value) < 80 and len(image_pages) < 6:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
            image_pages.append({"page": page_no, "mime_type": "image/png", "bytes": pix.tobytes("png")})
    return {"pages": pages, "images": image_pages}


def _call_gemini_contents(contents, is_json: bool = False, max_output_tokens: int = 12000):
    clients = get_gemini_clients()
    if not clients:
        return None
    for client in clients:
        for model_name in QUIZ_MODELS:
            try:
                config_kwargs = {"max_output_tokens": max_output_tokens}
                if is_json:
                    config_kwargs["response_mime_type"] = "application/json"
                if model_name.startswith("gemini-3."):
                    config_kwargs["thinking_config"] = types.ThinkingConfig(thinking_level="high")
                else:
                    config_kwargs["thinking_config"] = types.ThinkingConfig(thinking_budget=0, include_thoughts=False)
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents,
                    config=types.GenerateContentConfig(**config_kwargs),
                )
                if response and response.text:
                    return response.text
            except Exception:
                continue
    return None


def _vision_describe(image_bytes: bytes, mime_type: str, source_label: str) -> str:
    prompt = f"""
Anda adalah Vision Reader RoboMANTAP.
Sumber visual: {source_label}
Analisis fakta yang benar-benar terlihat pada gambar. Baca teks, angka, label, tabel, diagram, grafik, rumus, dan struktur visual yang relevan. Jangan mengarang bagian yang tidak terlihat. Gunakan Bahasa Indonesia yang rapi dan detail yang cukup untuk menjadi sumber soal serta media ajar.
"""
    try:
        part = types.Part.from_bytes(data=image_bytes, mime_type=mime_type)
        response = _call_gemini_contents([prompt, part], is_json=False, max_output_tokens=5000)
        return _trim_text(response or "", 8000)
    except Exception as exc:
        print(f"Vision material gagal: {exc}")
        return ""


def build_material_knowledge_pack(uploaded_files) -> dict:
    files = []
    for item in uploaded_files or []:
        name = str(getattr(item, "name", "material"))
        ext = os.path.splitext(name)[1].lower()
        if ext not in SUPPORTED_MATERIAL_EXTENSIONS:
            continue
        raw = item.getvalue() if hasattr(item, "getvalue") else bytes(item)
        record = {
            "name": name,
            "extension": ext,
            "mime_type": SUPPORTED_MATERIAL_EXTENSIONS[ext],
            "size": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest(),
            "content": "",
            "locator_map": [],
            "visual_analysis": [],
        }
        try:
            if ext == ".docx":
                record["content"] = _extract_docx_content(raw)
                record["locator_map"] = [{"type": "document", "locator": "paragraphs/tables"}]
            elif ext in {".pptx", ".ppt"}:
                pptx_bytes = raw if ext == ".pptx" else _convert_legacy_ppt_to_pptx(raw)
                if not pptx_bytes:
                    record["content"] = "Format .ppt lama diterima, tetapi runtime tidak menyediakan konverter LibreOffice. Gunakan .pptx untuk ekstraksi otomatis."
                else:
                    parsed = _extract_pptx_content(pptx_bytes)
                    blocks = []
                    for slide in parsed["slides"]:
                        block = f"[SLIDE {slide['slide']}]\n{slide['text']}"
                        if slide.get("notes"):
                            block += f"\n[CATATAN PEMBICARA]\n{slide['notes']}"
                        blocks.append(block)
                        record["locator_map"].append({"type": "slide", "locator": slide["slide"]})
                    record["content"] = _trim_text("\n\n".join(blocks), 50000)
                    for image in parsed["images"][:6]:
                        analysis = _vision_describe(image["bytes"], image["mime_type"], f"{name}, slide {image['slide']}")
                        if analysis:
                            record["visual_analysis"].append({"locator": f"slide {image['slide']}", "analysis": analysis})
            elif ext == ".pdf":
                parsed = _extract_pdf_content(raw)
                pages = []
                for page in parsed["pages"]:
                    pages.append(f"[HALAMAN {page['page']}]\n{page['text']}")
                    record["locator_map"].append({"type": "page", "locator": page["page"]})
                record["content"] = _trim_text("\n\n".join(pages), 50000)
                for image in parsed["images"][:6]:
                    analysis = _vision_describe(image["bytes"], image["mime_type"], f"{name}, halaman {image['page']}")
                    if analysis:
                        record["visual_analysis"].append({"locator": f"halaman {image['page']}", "analysis": analysis})
            else:
                analysis = _vision_describe(raw, record["mime_type"], name)
                record["content"] = analysis
                if analysis:
                    record["visual_analysis"].append({"locator": "image", "analysis": analysis})
        except Exception as exc:
            record["content"] = f"Ekstraksi gagal untuk {name}: {exc}"
        files.append(record)

    source_blocks = []
    for record in files:
        source_blocks.append(
            f"### FILE: {record['name']}\n"
            f"Format: {record['extension']}\n"
            f"SHA-256: {record['sha256']}\n"
            f"Konten:\n{record.get('content','')}"
        )
        for visual in record.get("visual_analysis", []):
            source_blocks.append(f"Visual {visual['locator']}:\n{visual['analysis']}")
    return {
        "files": files,
        "file_count": len(files),
        "source_text": _trim_text("\n\n".join(source_blocks), 110000),
    }


def _source_context_for_prompt(source_pack: dict | None, max_chars: int = 90000) -> str:
    if not source_pack:
        return "Tidak ada sumber materi terlampir."
    return _trim_text(source_pack.get("source_text", ""), max_chars)
